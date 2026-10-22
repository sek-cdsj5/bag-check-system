import streamlit as st
import pandas as pd
import openpyxl
import json
import os
from datetime import datetime
import plotly.express as px

# --- 頁面基本設定 ---
st.set_page_config(
    page_title="26-27 學年 學生書包秤重 Web 登記系統",
    page_icon="🎒",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- 自訂 CSS 樣式 ---
st.markdown("""
<style>
    .main-header {
        font-size: 1.8rem;
        font-weight: 700;
        color: #1E3A8A;
        margin-bottom: 0.5rem;
    }
    .sub-header {
        font-size: 1.0rem;
        color: #4B5563;
        margin-bottom: 1.5rem;
    }
    .card-pass {
        background-color: #DEF7EC;
        border-left: 5px solid #0E9F6E;
        padding: 10px 15px;
        border-radius: 6px;
        margin-bottom: 10px;
    }
    .card-over {
        background-color: #FDE8E8;
        border-left: 5px solid #F05252;
        padding: 10px 15px;
        border-radius: 6px;
        margin-bottom: 10px;
    }
    .login-box {
        max-width: 500px;
        margin: 50px auto;
        padding: 30px;
        background-color: #F9FAFB;
        border-radius: 12px;
        border: 1px solid #E5E7EB;
        box-shadow: 0 4px 6px -1px rgba(0,0,0,0.1);
    }
</style>
""", unsafe_allow_html=True)

# ==========================================
# 🔒 校內電郵身份驗證 (限制 @cdsj5.edu.mo)
# ==========================================
ALLOWED_DOMAIN = "@cdsj5.edu.mo"

query_params = st.query_params
saved_email = query_params.get("login_email", "")

if "user_email" not in st.session_state:
    st.session_state.user_email = saved_email if saved_email else ""

if not st.session_state.user_email:
    st.markdown('<div class="login-box">', unsafe_allow_html=True)
    st.markdown("## 🎒 學生書包秤重 Web 登記系統")
    st.markdown("#### 🔒 教職員身份驗證")
    st.info(f"本系統僅供校內教職員使用，請輸入貴校官方電郵 (`{ALLOWED_DOMAIN}`) 以進入系統。")

    input_email = st.text_input("請輸入學校 Email 帳號：", placeholder="username@cdsj5.edu.mo")

    if st.button("🔑 登入系統", type="primary", use_container_width=True):
        clean_email = input_email.strip().lower()
        if not clean_email:
            st.warning("⚠️ 請輸入 Email 帳號！")
        elif not clean_email.endswith(ALLOWED_DOMAIN):
            st.error(f"❌ 登入失敗：電郵網域不正確！必須以 `{ALLOWED_DOMAIN}` 結尾。")
        else:
            st.session_state.user_email = clean_email
            st.query_params["login_email"] = clean_email
            st.success("✅ 驗證成功，正在進入系統...")
            st.rerun()
            
    st.markdown('</div>', unsafe_allow_html=True)
    st.stop()

st.sidebar.write(f"👤 **目前登入者**：\n`{st.session_state.user_email}`")
if st.sidebar.button("🚪 登出系統"):
    st.session_state.user_email = ""
    if "login_email" in st.query_params:
        del st.query_params["login_email"]
    st.rerun()
st.sidebar.markdown("---")

OVERWEIGHT_REASONS = [
    "1. 書包材質偏重（非拖拉式書包）",
    "2. 水樽/壺內盛水很重",
    "3. 書本數量多",
    "4. 作業本數量多",
    "5. 文具及其他物品的數量多",
    "6. 其他，須註明"
]

CLASSES = [
    'P1A', 'P1B', 'P1C',
    'P2A', 'P2B', 'P2C',
    'P3A', 'P3B', 'P3C', 'P3D',
    'P4A', 'P4B', 'P4C',
    'P5A', 'P5B', 'P5C',
    'P6A', 'P6B', 'P6C', 'P6D'
]

DB_FILE = "survey_records.json"
XLSX_PATH = "學生書包秤重登記表_26-27學年_v5.xlsx"
DSEDJ_TEMPLATE_PATH = "學生書包重量抽查結果紀錄表及分析表.xlsx"

@st.cache_data
def load_student_database():
    students = {}
    if os.path.exists(XLSX_PATH):
        wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
        if '💾 學生體重資料庫' in wb.sheetnames:
            ws = wb['💾 學生體重資料庫']
            for row in range(4, ws.max_row + 1):
                key = ws.cell(row=row, column=1).value
                c_cls = ws.cell(row=row, column=2).value
                c_seat = ws.cell(row=row, column=3).value
                c_name = ws.cell(row=row, column=4).value
                c_gender = ws.cell(row=row, column=5).value
                c_weight = ws.cell(row=row, column=8).value
                if key and c_cls and c_seat:
                    students[str(key)] = {
                        "class": str(c_cls),
                        "seat": int(c_seat),
                        "name": str(c_name or ""),
                        "gender": str(c_gender or ""),
                        "weight": float(c_weight) if c_weight else None
                    }
    return students

STUDENT_DB = load_student_database()

def load_records():
    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_records(records):
    os.makedirs(os.path.dirname(DB_FILE) if os.path.dirname(DB_FILE) else ".", exist_ok=True)
    with open(DB_FILE, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

if "records" not in st.session_state:
    st.session_state.records = load_records()

st.sidebar.title("🎒 秤重登記系統")

user_mode = st.sidebar.radio(
    "請選擇操作模式：",
    ["📱 抽查老師登記端", "🖥️ 教務處/管理員儀表板", "ℹ️ 系統說明與超重標準"]
)

st.sidebar.markdown("---")
st.sidebar.info(f"💡 **系統狀態**：線上連線中\n\n年度：2026-2027 學年\n\n對照庫學生數：{len(STUDENT_DB)} 人")

# ==========================================
# 模式 1: 📱 抽查老師登記端
# ==========================================
if user_mode == "📱 抽查老師登記端":
    st.markdown('<div class="main-header">📱 學生書包秤重 — 班級抽查登記表</div>', unsafe_allow_html=True)
    st.markdown('<div class="sub-header">請選取班別，直接於座號欄位輸入抽查學生的座號，系統將自動帶出姓名與體重（共需登記 11 位）。</div>', unsafe_allow_html=True)

    col_cls, col_status = st.columns([2, 3])
    with col_cls:
        selected_class = st.selectbox("選擇抽查班別：", CLASSES, index=0)

    existing_class_data = st.session_state.records.get(selected_class, {})

    with col_status:
        if existing_class_data:
            st.success(f"✅ {selected_class} 曾於 {existing_class_data.get('updated_at', '')} 由 {existing_class_data.get('updated_by', '')} 提交過數據，再次提交將覆蓋更新。")
        else:
            st.info(f"📌 {selected_class} 尚未提交數據。請完成 11 位同學秤重後提交。")

    st.markdown("---")
    st.subheader(f"📋 {selected_class} 抽查學生名單 (11 人)")

    form_data = []
    
    for i in range(1, 12):
        st.markdown(f"##### 👤 第 {i} 位學生")
        c1, c2, c3, c4, c5 = st.columns([1.2, 2, 2, 2.5, 3])

        prev_row = existing_class_data.get("students", [])
        saved_seat = int(prev_row[i-1]["seat"]) if (i <= len(prev_row) and prev_row[i-1].get("seat")) else None
        default_seat = saved_seat if (saved_seat and saved_seat > 0) else None
        
        raw_prev_bag = prev_row[i-1]["bag_weight"] if i <= len(prev_row) else None
        default_bag = float(raw_prev_bag) if raw_prev_bag and raw_prev_bag > 0 else None

        default_reason = prev_row[i-1]["reason"] if i <= len(prev_row) else OVERWEIGHT_REASONS[0]
        default_note = prev_row[i-1]["note"] if i <= len(prev_row) else ""

        with c1:
            seat_no = st.number_input(
                f"座號 #{i}", 
                min_value=0, 
                max_value=45, 
                value=default_seat, 
                step=1,
                placeholder="請輸入座號",
                key=f"seat_{selected_class}_{i}"
            )
        
        if seat_no and seat_no > 0:
            lookup_key = f"{selected_class}-{int(seat_no)}"
            s_info = STUDENT_DB.get(lookup_key, {})
            s_name = s_info.get("name", "⚠️ 查無此座號學生")
            s_weight = s_info.get("weight", None)
        else:
            s_name = "請輸入座號"
            s_weight = None

        with c2:
            st.text_input(f"姓名 #{i}", value=s_name, disabled=True, key=f"name_{selected_class}_{seat_no}_{i}")
        
        with c3:
            weight_str = f"{s_weight} kg" if s_weight is not None else "---"
            st.text_input(f"體重 #{i}", value=weight_str, disabled=True, key=f"weight_{selected_class}_{seat_no}_{i}")

        with c4:
            bag_w = st.number_input(
                f"書包重量 (kg) #{i}", 
                min_value=0.0, 
                max_value=20.0, 
                step=0.1, 
                value=default_bag, 
                placeholder="例如：2.5",
                key=f"bag_{selected_class}_{i}"
            )

        bag_weight_val = float(bag_w) if bag_w is not None else 0.0

        ratio = (bag_weight_val / s_weight * 100) if (s_weight and bag_weight_val > 0) else 0.0
        is_overweight = ratio > 15.0 if (s_weight and bag_weight_val > 0) else False

        with c5:
            if seat_no and seat_no > 0 and bag_weight_val > 0 and s_weight:
                if is_overweight:
                    st.markdown(f"<div class='card-over'>❌ <b>超重 {ratio:.1f}%</b> (>15%)</div>", unsafe_allow_html=True)
                else:
                    st.markdown(f"<div class='card-pass'>🟢 <b>達標 {ratio:.1f}%</b> (≤15%)</div>", unsafe_allow_html=True)
            else:
                st.caption("請輸入座號與重量")

        reason_val = ""
        note_val = ""
        if is_overweight:
            rc1, rc2 = st.columns([2, 2])
            with rc1:
                reason_idx = OVERWEIGHT_REASONS.index(default_reason) if default_reason in OVERWEIGHT_REASONS else 0
                reason_val = st.selectbox(f"超重原因 #{i}", OVERWEIGHT_REASONS, index=reason_idx, key=f"reason_{selected_class}_{i}")
            with rc2:
                note_val = st.text_input(f"備註說明 #{i}", value=default_note, placeholder="其他原因註明...", key=f"note_{selected_class}_{i}")

        form_data.append({
            "seq": i,
            "seat": int(seat_no) if seat_no else 0,
            "name": s_name,
            "body_weight": s_weight,
            "bag_weight": bag_weight_val,
            "ratio": round(ratio, 2) if ratio else 0.0,
            "status": "超重" if is_overweight else ("達標" if bag_weight_val > 0 else "未輸入"),
            "reason": reason_val if is_overweight else "",
            "note": note_val if is_overweight else "",
            "submitted_by": st.session_state.user_email
        })
        st.markdown("<hr style='margin: 5px 0;'>", unsafe_allow_html=True)

    st.markdown("### 📤 資料提交確認")
    valid_count = sum(1 for item in form_data if item["seat"] > 0 and item["bag_weight"] > 0)
    st.progress(valid_count / 11)
    st.write(f"目前完成有效輸入人數：**{valid_count} / 11** 人")

    if st.button("🚀 提交本班抽查紀錄", type="primary", use_container_width=True):
        if valid_count < 11:
            st.warning("⚠️ 提醒：目前尚未完成 11 位學生的完整輸入，確定要先提交嗎？")
        
        st.session_state.records[selected_class] = {
            "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "updated_by": st.session_state.user_email,
            "students": form_data
        }
        save_records(st.session_state.records)
        st.success(f"🎉 【{selected_class}】共 {valid_count} 筆資料已成功寫入中央資料庫！管理員儀表板已同步更新。")
        st.balloons()

# ==========================================
# 模式 2: 🖥️ 教務處/管理員儀表板
# ==========================================
elif user_mode == "🖥️ 教務處/管理員儀表板":
    st.markdown('<div class="main-header">🖥️ 全校學生書包秤重 — 管理員儀表板</div>', unsafe_allow_html=True)
    st.markdown('<div class="sub-header">即時監控全校 20 班抽查進度、分析超重原因，並可一鍵產出教青局官方標準格式總報告。</div>', unsafe_allow_html=True)

    records = st.session_state.records

    total_classes_submitted = len(records)
    all_students_flat = []
    for cls, cdata in records.items():
        for st_item in cdata.get("students", []):
            if st_item["seat"] > 0 and st_item["bag_weight"] > 0:
                item_copy = st_item.copy()
                item_copy["class"] = cls
                all_students_flat.append(item_copy)

    total_surveyed = len(all_students_flat)
    total_pass = sum(1 for s in all_students_flat if s["status"] == "達標")
    total_over = sum(1 for s in all_students_flat if s["status"] == "超重")
    pass_rate = (total_pass / total_surveyed * 100) if total_surveyed > 0 else 0.0

    underweight_count = sum(1 for s in all_students_flat if "體重過輕" in s.get("reason", ""))

    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("已提交班級", f"{total_classes_submitted} / 20 班")
    m2.metric("總抽查學生數", f"{total_surveyed} 人")
    m3.metric("全校達標率", f"{pass_rate:.1f} %")
    m4.metric("超重人數", f"{total_over} 人", delta_color="inverse")
    m5.metric("因『體重過輕』超重", f"{underweight_count} 人")

    st.markdown("---")

    c_left, c_right = st.columns([1, 1])

    with c_left:
        st.subheader("📊 超重原因分類統計")
        if total_over > 0:
            reasons_list = [s["reason"] for s in all_students_flat if s["status"] == "超重" and s["reason"]]
            df_reasons = pd.Series(reasons_list).value_counts().reset_index()
            df_reasons.columns = ["超重原因", "人數"]
            fig_reasons = px.pie(df_reasons, values="人數", names="超重原因", color_discrete_sequence=px.colors.qualitative.Pastel, hole=0.4)
            st.plotly_chart(fig_reasons, use_container_width=True)
        else:
            st.info("尚無超重數據紀錄。")

    with c_right:
        st.subheader("🏫 各班達標率與完成狀態")
        class_progress = []
        for cls in CLASSES:
            if cls in records:
                st_list = [s for s in records[cls]["students"] if s["seat"] > 0 and s["bag_weight"] > 0]
                c_pass = sum(1 for s in st_list if s["status"] == "達標")
                c_rate = (c_pass / len(st_list) * 100) if st_list else 0.0
                updater = records[cls].get("updated_by", "未知")
                class_progress.append({"班別": cls, "已填人數": len(st_list), "達標率(%)": round(c_rate, 1), "狀態": "✅ 已提交", "提交者": updater})
            else:
                class_progress.append({"班別": cls, "已填人數": 0, "達標率(%)": 0.0, "狀態": "⏳ 未提交", "提交者": "-"})

        df_progress = pd.DataFrame(class_progress)
        st.dataframe(df_progress, use_container_width=True, height=280)

    st.markdown("---")

    st.subheader("📥 匯出教青局官方標準格式報告")
    
    def generate_dsedj_report_bytes():
        if os.path.exists(DSEDJ_TEMPLATE_PATH):
            wb = openpyxl.load_workbook(DSEDJ_TEMPLATE_PATH)
        else:
            wb = openpyxl.Workbook()
            wb.active.title = "學生書包重量抽查結果紀錄表"
            wb.create_sheet(title="學生書包重量抽查結果分析表")

        grade_mapping = {
            "P1": "小一", "P2": "小二", "P3": "小三",
            "P4": "小四", "P5": "小五", "P6": "小六"
        }

        if "學生書包重量抽查結果紀錄表" in wb.sheetnames:
            ws_rec = wb["學生書包重量抽查結果紀錄表"]
            start_row = 7
            row_idx = 0
            
            for cls in CLASSES:
                prefix = cls[:2]
                grade_name = grade_mapping.get(prefix, cls)

                cdata = records.get(cls, {})
                s_list = cdata.get("students", [])
                for s in s_list:
                    if s["seat"] > 0 and s["bag_weight"] > 0:
                        curr_row = start_row + row_idx
                        ws_rec.cell(row=curr_row, column=1, value=row_idx + 1)
                        ws_rec.cell(row=curr_row, column=2, value=grade_name)
                        ws_rec.cell(row=curr_row, column=3, value=s["body_weight"])
                        ws_rec.cell(row=curr_row, column=4, value=s["bag_weight"])
                        ws_rec.cell(row=curr_row, column=5, value=s["ratio"] / 100.0 if s["ratio"] else 0)
                        ws_rec.cell(row=curr_row, column=6, value="ü" if s["status"] == "達標" else "û")
                        
                        # 處理超重原因代號與備註
                        reason_text = s.get("reason", "")
                        note_text = s.get("note", "").strip()
                        
                        # 提取代號（例如從 "1. 書包材質偏重..." 提取出數字）
                        reason_code = reason_text.split(".")[0].strip() if "." in reason_text else reason_text[:1]
                        
                        if note_text:
                            final_reason_val = f"{reason_code}，{note_text}"
                        else:
                            final_reason_val = reason_code

                        ws_rec.cell(row=curr_row, column=7, value=final_reason_val)
                        
                        row_idx += 1

        if "學生書包重量抽查結果分析表" in wb.sheetnames:
            ws_ana = wb["學生書包重量抽查結果分析表"]
            ws_ana.cell(row=4, column=5, value=len(STUDENT_DB))
            ws_ana.cell(row=5, column=3, value=total_surveyed)
            ws_ana.cell(row=6, column=3, value=total_over)
            ws_ana.cell(row=7, column=3, value=total_pass)

        out_path = "dsedj_report_export.xlsx"
        wb.save(out_path)
        with open(out_path, "rb") as f:
            return f.read()

    dsedj_excel_bytes = generate_dsedj_report_bytes()
    st.download_button(
        label="📄 下載教青局官方標準格式總報告 (.xlsx)",
        data=dsedj_excel_bytes,
        file_name=f"教青局學生書包重量抽查結果紀錄及分析表_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )

    with st.expander("🔍 檢視全校已提交之詳細學生明細數據"):
        if all_students_flat:
            df_all = pd.DataFrame(all_students_flat)
            st.dataframe(df_all[["class", "seat", "name", "body_weight", "bag_weight", "ratio", "status", "reason", "note", "submitted_by"]], use_container_width=True)
        else:
            st.info("尚無學生明細數據。")

    # ==========================================
    # 🗑️ 管理員資料清理專區
    # ==========================================
    st.markdown("---")
    st.subheader("⚙️ 系統資料管理與清除")
    with st.expander("⚠️ 危險操作區：清除已登記的資料"):
        st.warning("注意：此操作將會清空所有已提交的班級抽查紀錄（`survey_records.json`），且無法復原！請在匯出備份好報告後才進行此操作。")
        confirm_delete = st.checkbox("我確定要清除全校所有已登記的抽查資料")
        if st.button("🗑️ 立即清空所有登記資料", type="secondary"):
            if confirm_delete:
                st.session_state.records = {}
                if os.path.exists(DB_FILE):
                    try:
                        os.remove(DB_FILE)
                    except Exception:
                        pass
                st.success("✅ 已成功清空所有已登記的抽查資料！")
                st.rerun()
            else:
                st.error("❌ 請先勾選「我確定要清除全校所有已登記的抽查資料」才能執行清除。")

# ==========================================
# 模式 3: ℹ️ 系統說明與超重標準
# ==========================================
else:
    st.markdown('<div class="main-header">ℹ️ 學生書包秤重系統 — 評定標準與說明</div>', unsafe_allow_html=True)
    st.markdown("""
    ### 🎯 評定標準與公式
    * **運算公式**：書包佔體重百分比 = (書包重量 kg / 學生體重 kg) * 100%
    * **🟢 達標標準**：書包重量佔體重百分比 <= 15%。
    * **🔴 超重標準**：書包重量佔體重百分比 > 15%。

    ---

    ### 📝 超重原因分類
    1. **1. 書包材質偏重（非拖拉式書包）**
    2. **2. 水樽/壺內盛水很重**
    3. **3. 書本數量多**
    4. **4. 作業本數量多**
    5. **5. 文具及其他物品的數量多**
    6. **6. 其他，須註明**

    ---

    ### 📱 網頁版操作流程
    1. 輸入學校官方電郵 (`@cdsj5.edu.mo`) 完成身份驗證。
    2. 切換至 **「📱 抽查老師登記端」**。
    3. 下拉選取貴班班別（P1A ~ P6D）。
    4. 於座號欄位直接輸入抽查學生的座號，系統會**即時自動帶出學生姓名與體重**。
    5. 輸入書包重量，系統自動判斷是否超重。
    6. 完成後點擊 **「提交本班抽查紀錄」**。
    7. 教務處可在 **「🖥️ 教務處/管理員儀表板」** 查看統計並一鍵下載教青局官方標準格式報告（若填有備註，將自動與代號以「，」逗號分隔合併顯示）。
    """)
